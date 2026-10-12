"""Build the lab presentation for Zou et al., Lancet Digit Health 2026 (agentic AI in eye care)."""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).parent
IMG = HERE / "img"
OUT = HERE / "eye_agentic_ai_presentation.pptx"

NAVY = RGBColor(0x1F, 0x2A, 0x5A)
ACCENT = RGBColor(0xC8, 0x1E, 0x2B)
TEAL = RGBColor(0x2E, 0x7D, 0x9A)
GREEN = RGBColor(0x4C, 0x9A, 0x4A)
ORANGE = RGBColor(0xF0, 0xA2, 0x1C)
LIGHT = RGBColor(0xF2, 0xF4, 0xF8)
GREY = RGBColor(0x55, 0x5B, 0x6E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
TOTAL = 15


def box(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.5)
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, content, size=18, color=GREY, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
    """content: str or list of str / (str, dict) for per-paragraph overrides."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    items = content if isinstance(content, list) else [content]
    for i, item in enumerate(items):
        t, opt = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opt.get("align", align)
        p.space_after = Pt(opt.get("space", 6))
        r = p.add_run()
        r.text = t
        f = r.font
        f.name = FONT
        f.size = Pt(opt.get("size", size))
        f.bold = opt.get("bold", bold)
        f.italic = opt.get("italic", italic)
        f.color.rgb = opt.get("color", color)
    return tb


def bullets(slide, x, y, w, h, items, size=18, color=GREY):
    """items: list of str or (str, level)."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        t, lvl = (item, 0) if isinstance(item, str) else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8 if lvl == 0 else 4)
        r = p.add_run()
        r.text = ("\u25A0  " if lvl == 0 else "\u2013  ") + t
        r.font.name = FONT
        r.font.size = Pt(size if lvl == 0 else size - 2)
        r.font.color.rgb = color
        if lvl:
            p.level = 1
    return tb


def header(slide, title, n, subtitle=None):
    box(slide, 0, 0, SW, Inches(1.0), fill=NAVY)
    box(slide, 0, Inches(1.0), SW, Inches(0.06), fill=ACCENT)
    text(slide, Inches(0.5), Inches(0.12), Inches(11.5), Inches(0.8), title,
         size=30, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    if subtitle:
        text(slide, Inches(0.5), Inches(1.12), Inches(12.3), Inches(0.45), subtitle,
             size=15, color=TEAL, italic=True)
    text(slide, Inches(0.5), SH - Inches(0.45), Inches(9), Inches(0.35),
         "Zou et al., Lancet Digital Health 2026; 8: 100967", size=11, color=GREY)
    text(slide, SW - Inches(1.5), SH - Inches(0.45), Inches(1.0), Inches(0.35),
         f"{n} / {TOTAL}", size=11, color=GREY, align=PP_ALIGN.RIGHT)


def picture_fit(slide, path, x, y, max_w, max_h, border=True):
    w_px, h_px = Image.open(path).size
    scale = min(max_w / w_px, max_h / h_px)
    w, h = int(w_px * scale), int(h_px * scale)
    px, py = x + (max_w - w) // 2, y + (max_h - h) // 2
    pic = slide.shapes.add_picture(str(path), px, py, w, h)
    if border:
        pic.line.color.rgb = RGBColor(0xC8, 0xCC, 0xD6)
        pic.line.width = Pt(1)
    return pic


def card(slide, x, y, w, h, title, body, color, title_size=17, body_size=14):
    box(slide, x, y, w, h, fill=LIGHT)
    box(slide, x, y, Inches(0.09), h, fill=color)
    text(slide, x + Inches(0.2), y + Inches(0.08), w - Inches(0.3), Inches(0.5),
         title, size=title_size, color=color, bold=True)
    text(slide, x + Inches(0.2), y + Inches(0.55), w - Inches(0.3), h - Inches(0.6),
         body, size=body_size, color=GREY)


# 1. Title -------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, SW, SH, fill=NAVY)
box(s, 0, Inches(4.55), SW, Inches(0.07), fill=ACCENT)
text(s, Inches(0.8), Inches(0.9), Inches(11.7), Inches(0.5), "LAB PAPER PRESENTATION",
     size=16, color=ORANGE, bold=True)
text(s, Inches(0.8), Inches(1.5), Inches(11.7), Inches(2.0),
     "Agentic Artificial Intelligence in Eye Care:\nIs Clinical Autonomy Finally Within Reach?",
     size=40, color=WHITE, bold=True)
text(s, Inches(0.8), Inches(3.55), Inches(11.7), Inches(0.9),
     "Ke Zou, Jocelyn Hui Lin Goh, Gabriel Dawei Yang, Sahana Srinivasan, Qingyu Chen, "
     "Pearse A. Keane, Yih-Chung Tham*",
     size=17, color=RGBColor(0xD6, 0xDB, 0xEA))
text(s, Inches(0.8), Inches(4.85), Inches(11.7), Inches(1.6), [
    ("The Lancet Digital Health, Vol 8, March 2026, 100967  (Comment)", {"size": 18, "color": WHITE, "bold": True}),
    ("Published online 19 Feb 2026   |   doi: 10.1016/j.landig.2025.100967", {"size": 15, "color": RGBColor(0xD6, 0xDB, 0xEA)}),
    ("National University of Singapore  \u2022  Singapore Eye Research Institute  \u2022  Yale  \u2022  Moorfields / UCL",
     {"size": 15, "color": RGBColor(0xD6, 0xDB, 0xEA)}),
])
text(s, Inches(0.8), Inches(6.3), Inches(11.7), Inches(0.9), [
    ("Presented by: Sheikh Sayed Bin Rahman", {"size": 17, "color": ORANGE, "bold": True, "space": 2}),
    ("PIC Lab, Kumoh National Institute of Technology", {"size": 15, "color": RGBColor(0xD6, 0xDB, 0xEA)}),
])

# 2. Journal -----------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "About the Journal: The Lancet Digital Health", 2,
       "A top-tier Q1 journal at the intersection of medicine and digital technology")
metrics = [
    ("SCImago (SJR)", "6.838", NAVY,
     "Q1 in every indexed category:\n\u2022 Health Informatics\n\u2022 Health Information Management\n"
     "\u2022 Medicine (miscellaneous)\n\u2022 Decision Sciences (misc.)\nLatest SJR edition: Q1 (2025)"),
    ("Clarivate JCR (SCIE)", "IF 24.1 \u2192 25.5", ACCENT,
     "Journal Impact Factor:\n\u2022 24.1 (2024)\n\u2022 25.5 (2025)\n\nTop ~2% of its categories\n"
     "98.1st percentile \u2014 comfortably Q1"),
    ("Scopus CiteScore", "38.4", TEAL,
     "CiteScore 38.4\n\nRanked Q1 in Scopus"),
]
cw, ch, gap = Inches(3.9), Inches(4.6), Inches(0.3)
x0 = (SW - 3 * cw - 2 * gap) // 2
for i, (name, val, col, body) in enumerate(metrics):
    x = x0 + i * (cw + gap)
    y = Inches(1.75)
    box(s, x, y, cw, ch, fill=LIGHT)
    box(s, x, y, cw, Inches(0.6), fill=col)
    text(s, x, y, cw, Inches(0.6), name, size=19, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x, y + Inches(0.75), cw, Inches(0.9), val, size=38, color=col, bold=True,
         align=PP_ALIGN.CENTER)
    q = box(s, x + cw // 2 - Inches(0.45), y + Inches(1.65), Inches(0.9), Inches(0.42),
            fill=GREEN, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x + cw // 2 - Inches(0.45), y + Inches(1.65), Inches(0.9), Inches(0.42), "Q1",
         size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + Inches(0.3), y + Inches(2.2), cw - Inches(0.6), ch - Inches(2.3), body,
         size=15, color=GREY)
text(s, Inches(0.7), Inches(6.5), Inches(12), Inches(0.5),
     "Article type: Comment (perspective / position paper) \u2014 proposes a framework rather than reporting an experiment.",
     size=14, color=NAVY, italic=True, align=PP_ALIGN.CENTER)

# 3. Motivation --------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Background & Motivation", 3, "Why current foundation models are not enough for the clinic")
bullets(s, Inches(0.6), Inches(1.75), Inches(6.4), Inches(5), [
    "Foundation models \u2014 LLMs and large vision-language models (VLMs) \u2014 drive major progress in medical AI",
    "But they are reactive: they only produce output in response to an explicit prompt",
    "Constrained by task-specific training; little adaptability or autonomy",
    "Cannot continuously perceive changes in clinical context, integrate new information, adjust reasoning, "
    "or initiate actions without a human",
], size=18)
card(s, Inches(7.4), Inches(1.75), Inches(5.4), Inches(2.25), "The clinical problem",
     "Global shortage of ophthalmologists + ageing populations + comorbidities "
     "\u2192 rising workload and demand for personalised care.", ACCENT, body_size=16)
card(s, Inches(7.4), Inches(4.2), Inches(5.4), Inches(2.25), "The paper's question",
     "Can agentic AI \u2014 autonomous, adaptive, collaborative agents \u2014 take over routine "
     "clinical tasks in eye care so clinicians focus on final decisions?", TEAL, body_size=16)

# 4. Concepts ----------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Key Concepts: From Foundation Models to Agentic AI", 4)
stages = [
    ("Foundation model", "LLM / VLM\n\nReactive: prompt \u2192 output\nTask-specific", GREY),
    ("AI agent", "Foundation model\n+ goal-oriented reasoning\n+ tool use", TEAL),
    ("Agentic AI\n(multi-agent)", "Multiple specialised agents\nshare context, interact &\norchestrate tasks dynamically", NAVY),
]
sw_, sh_ = Inches(3.6), Inches(2.6)
gap = Inches(0.75)
x0 = (SW - 3 * sw_ - 2 * gap) // 2
for i, (t, b, col) in enumerate(stages):
    x = x0 + i * (sw_ + gap)
    y = Inches(1.5)
    box(s, x, y, sw_, Inches(0.85), fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x, y, sw_, Inches(0.85), t, size=19, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    box(s, x, y + Inches(0.95), sw_, sh_ - Inches(0.95), fill=LIGHT)
    text(s, x + Inches(0.15), y + Inches(1.05), sw_ - Inches(0.3), sh_ - Inches(1.1), b,
         size=15, color=GREY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if i < 2:
        box(s, x + sw_ + Inches(0.12), y + Inches(1.0), gap - Inches(0.24), Inches(0.6),
            fill=ACCENT, shape=MSO_SHAPE.RIGHT_ARROW)
text(s, Inches(0.6), Inches(4.4), Inches(12), Inches(0.5),
     "Key attributes of agentic AI in medicine", size=19, color=NAVY, bold=True)
attrs = ["Real-time adaptability", "Contextual awareness", "Autonomous decision making",
         "Planning & tool selection", "Proactive action (e.g., adjust treatment plans)"]
aw = Inches(2.35)
for i, a in enumerate(attrs):
    x = Inches(0.6) + i * (aw + Inches(0.12))
    box(s, x, Inches(5.0), aw, Inches(1.1), fill=WHITE, line=TEAL, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x, Inches(5.0), aw, Inches(1.1), a, size=14, color=NAVY, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, Inches(0.6), Inches(6.3), Inches(12), Inches(0.5),
     "Goes beyond passive prediction: integrates history, symptoms and current condition to suggest diagnoses "
     "and personalised management.", size=14, italic=True)

# 5. Methodology overview ----------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Methodology Overview: The Proposed Framework", 5,
       "Figure: Shift from traditional AI to agentic AI in medical workflows (extracted from the paper)")
picture_fit(s, IMG / "figure_full.png", Inches(0.4), Inches(1.6), Inches(8.3), Inches(5.35))
card(s, Inches(8.95), Inches(1.6), Inches(4.0), Inches(1.6), "(A) Traditional AI",
     "AI assists diagnosis only on clinician request; heavy oversight.", GREY, body_size=14)
card(s, Inches(8.95), Inches(3.35), Inches(4.0), Inches(1.6), "(B) Medical agentic AI",
     "Agents plan, order tests, diagnose, and report; clinician makes final decision.", TEAL, body_size=14)
card(s, Inches(8.95), Inches(5.1), Inches(4.0), Inches(1.85), "(C) Eye agentic system",
     "Cataract use case: four specialised agents orchestrate the full consultation.", ACCENT, body_size=14)

# 6. Panel A -----------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Methodology (A): Traditional AI-Assisted Care Model", 6, "The baseline the paper argues against")
picture_fit(s, IMG / "panel_A.png", Inches(0.4), Inches(1.65), Inches(12.5), Inches(2.6))
card(s, Inches(0.5), Inches(4.5), Inches(6.0), Inches(2.4), "Workflow",
     "Patient history \u2192 clinician orders diagnostic tests \u2192 information (history, examinations, "
     "records) collected in the EHR \u2192 AI model gives a preliminary diagnosis \u2192 clinician "
     "provides diagnosis & treatment plan.", NAVY, body_size=15)
card(s, Inches(6.8), Inches(4.5), Inches(6.0), Inches(2.4), "Limitations",
     "\u2022 AI acts only when the clinician requests it\n\u2022 Limited autonomy, single isolated task\n"
     "\u2022 Requires substantial clinician oversight\n\u2022 Time consuming for the clinician", ACCENT, body_size=15)

# 7. Panel B -----------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Methodology (B): Care Model with Medical Agentic AI", 7,
       "Agents act autonomously along the patient journey; clinician focuses on final decisions")
picture_fit(s, IMG / "panel_B.png", Inches(0.4), Inches(1.65), Inches(12.5), Inches(3.0))
loop = [("1. Devise clinical plan", NAVY), ("2. Order tests", TEAL),
        ("3. Preliminary diagnosis", GREEN), ("4. Generate reports", ACCENT)]
lw = Inches(2.7)
for i, (t, col) in enumerate(loop):
    x = Inches(0.6) + i * (lw + Inches(0.42))
    box(s, x, Inches(4.9), lw, Inches(0.7), fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x, Inches(4.9), lw, Inches(0.7), t, size=16, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if i < 3:
        box(s, x + lw + Inches(0.06), Inches(5.07), Inches(0.3), Inches(0.36), fill=GREY,
            shape=MSO_SHAPE.RIGHT_ARROW)
bullets(s, Inches(0.6), Inches(5.85), Inches(12.2), Inches(1.2), [
    "Multi-agent system takes the patient history, then handles test ordering, treatment planning, report drafting and preliminary interpretation",
    "Routine / cognitively demanding work (history taking, note summarisation) is offloaded, which reduces workload and improves efficiency",
], size=15)

# 8. Why ophthalmology -------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Why Ophthalmology Is a Good Fit for Agentic AI", 8)
reasons = [
    ("Workforce shortage", "Global shortage of ophthalmologists vs. growing service demand.", ACCENT),
    ("Rising complexity", "Ageing populations and comorbid conditions demand personalised, comprehensive care.", ORANGE),
    ("Multimodal data", "Rich imaging (fundus, slit-lamp, OCT) + clinical measurements \u2014 suited to VLMs + LLMs.", TEAL),
    ("Protocol-driven", "Standardised exam sequences (VA, IOP, refraction, dilation) map naturally onto agent tasks.", NAVY),
    ("Bridging specialties", "An eye-specific agent with ophthalmic + systemic knowledge can bridge subspecialty gaps.", GREEN),
    ("Resource optimisation", "Pre-completes consultation steps before the clinician's final review.", GREY),
]
cw, ch = Inches(3.95), Inches(2.4)
for i, (t, b, col) in enumerate(reasons):
    r, c = divmod(i, 3)
    card(s, Inches(0.5) + c * (cw + Inches(0.2)), Inches(1.5) + r * (ch + Inches(0.25)), cw, ch,
         t, b, col, title_size=19, body_size=16)

# 9. Panel C -----------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Methodology (C): Eye Agentic AI System for Cataract", 9,
       "Use case: patient needs \u2192 four collaborating agents \u2192 ophthalmologist")
picture_fit(s, IMG / "panel_C.png", Inches(0.3), Inches(1.65), Inches(12.7), Inches(3.6))
card(s, Inches(0.5), Inches(5.45), Inches(3.9), Inches(1.5), "Input: patient's needs",
     "Initial consultation, surgical & medication queries, chronic conditions", ORANGE, body_size=14)
card(s, Inches(4.7), Inches(5.45), Inches(3.9), Inches(1.5), "Eye agentic AI system",
     "Imaging \u2192 Diagnosis \u2192 Summarisation \u2192 Scheduling agents", ACCENT, body_size=14)
card(s, Inches(8.9), Inches(5.45), Inches(3.9), Inches(1.5), "Output: ophthalmologist",
     "Confirms diagnosis, administers treatment, follow-up examination", GREEN, body_size=14)

# 10. Agent pipeline ---------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Methodology (C): Agent-by-Agent Pipeline", 10,
       "Each agent owns one subtask; together they complete the consultation before clinician review")
agents = [
    ("Imaging agent", "LLM", NAVY,
     "Gathers eye complaints & visual demands; autonomously orders tests: VA, intraocular pressure, "
     "refraction, dilation, fundus photo"),
    ("Diagnosis agent", "VLM", TEAL,
     "Analyses acquired ocular images; preliminary diagnosis: (1) cataract, (2) other comorbidities. "
     "Evidence: lens opacity, reduced VA"),
    ("Summarisation agent", "LLM", GREEN,
     "Compiles diagnostic information into summarised reports and a preliminary treatment plan; "
     "forwards to clinician"),
    ("Scheduling agent", "Triggered", ACCENT,
     "If significant cataract: books surgery, recommends a suitable artificial lens (IOL), "
     "plans tailored post-op follow-up"),
]
aw, ah = Inches(2.85), Inches(4.3)
gap = Inches(0.35)
x0 = (SW - 4 * aw - 3 * gap) // 2
for i, (t, m, col, b) in enumerate(agents):
    x, y = x0 + i * (aw + gap), Inches(1.75)
    box(s, x, y, aw, ah, fill=LIGHT)
    box(s, x, y, aw, Inches(0.75), fill=col)
    text(s, x, y, aw, Inches(0.75), t, size=18, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    box(s, x + aw // 2 - Inches(0.7), y + Inches(0.9), Inches(1.4), Inches(0.4), fill=WHITE, line=col,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x + aw // 2 - Inches(0.7), y + Inches(0.9), Inches(1.4), Inches(0.4), m, size=13,
         color=col, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + Inches(0.15), y + Inches(1.45), aw - Inches(0.3), ah - Inches(1.5), b, size=15)
    if i < 3:
        box(s, x + aw + Inches(0.04), y + Inches(2.1), gap - Inches(0.08), Inches(0.45), fill=GREY,
            shape=MSO_SHAPE.RIGHT_ARROW)
text(s, Inches(0.6), Inches(6.25), Inches(12.1), Inches(0.7),
     "Result: consultation, diagnostic interpretation and follow-up preparation are done in advance \u2192 "
     "streamlined, patient-centred workflow; the clinician retains the final decision.",
     size=15, color=NAVY, bold=True, align=PP_ALIGN.CENTER)

# 11. Challenges -------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Challenges of Agentic AI in Healthcare", 11, "Five real-world barriers identified by the authors")
chal = [
    ("1. Computational cost", "Storing/processing clinical history, imaging and inter-agent dialogue strains edge and mobile devices."),
    ("2. Fragile coordination", "Without standard protocols and shared goals, agents may give inconsistent outputs or compete for resources (e.g., emergency triage)."),
    ("3. Emergent behaviour", "Complex interactions can be unsafe, e.g., scheduling and diagnostic agents mis-prioritising high-risk patients."),
    ("4. Debugging & traceability", "Tracing decisions across asynchronous, interdependent agents is hard and erodes clinician trust."),
    ("5. Fairness, ethics & safety", "Remain considerably under-explored for agentic deployments."),
]
for i, (t, b) in enumerate(chal):
    y = Inches(1.65) + i * Inches(1.05)
    box(s, Inches(0.6), y, Inches(3.6), Inches(0.9), fill=ACCENT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, Inches(0.6), y, Inches(3.6), Inches(0.9), t, size=17, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    box(s, Inches(4.35), y, Inches(8.4), Inches(0.9), fill=LIGHT)
    text(s, Inches(4.5), y, Inches(8.2), Inches(0.9), b, size=15, anchor=MSO_ANCHOR.MIDDLE)

# 12. Solutions --------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Proposed Solutions & Key Enablers", 12, "Coordinated advances needed across several fronts")
sols = [
    ("Human-in-the-loop", "Expert oversight to guide agents and correct errors in time.", NAVY),
    ("Standardised protocols & semantic planning", "Better coordination, less ambiguity, outputs aligned with clinical workflow & intent.", TEAL),
    ("Verifiable governance", "Real-time behaviour monitoring, enforced safety limits, models predicting agent\u2013environment interaction.", ACCENT),
    ("Cross-specialty reasoning", "Shared causal reasoning + coherent context modelling between ophthalmic and systemic agents for interpretability.", GREEN),
]
cw, ch = Inches(6.0), Inches(2.0)
for i, (t, b, col) in enumerate(sols):
    r, c = divmod(i, 2)
    card(s, Inches(0.5) + c * (cw + Inches(0.3)), Inches(1.65) + r * (ch + Inches(0.25)), cw, ch,
         t, b, col, title_size=19, body_size=16)
box(s, Inches(0.5), Inches(6.2), Inches(12.3), Inches(0.75), fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, Inches(0.5), Inches(6.2), Inches(12.3), Inches(0.75),
     "Privacy and fairness must be safeguarded at the core to earn clinician & patient trust and ensure equitable care.",
     size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# 13. Conclusion -------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Conclusion", 13)
bullets(s, Inches(0.6), Inches(1.5), Inches(12.2), Inches(4.0), [
    "Agentic AI moves medical AI from reactive prediction to autonomous, interactive and adaptive assistance",
    "Can analyse clinical patterns, assist diagnosis, personalise treatment and support longitudinal care",
    "Ophthalmology is a compelling first domain: multimodal and protocol-driven",
    "The cataract case shows four agents completing the consultation pipeline before clinician review",
    "Clinical autonomy is within reach, but only with human oversight, standard protocols, verifiable governance, and privacy/fairness by design",
], size=19)
box(s, Inches(0.6), Inches(5.55), Inches(12.2), Inches(1.3), fill=LIGHT)
box(s, Inches(0.6), Inches(5.55), Inches(0.09), Inches(1.3), fill=ACCENT)
text(s, Inches(0.85), Inches(5.6), Inches(11.8), Inches(1.2),
     "Answer to the title question: \u201cpartially\u201d \u2014 the technology is ready for supervised "
     "autonomy on routine tasks; full autonomy awaits solutions to safety, coordination and trust.",
     size=17, color=NAVY, italic=True, anchor=MSO_ANCHOR.MIDDLE)

# 14. Critical discussion ----------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Critical Discussion", 14, "Our assessment for lab discussion")
card(s, Inches(0.5), Inches(1.65), Inches(6.0), Inches(4.4), "Strengths",
     "\u2022 Clear conceptual model contrasting traditional vs. agentic AI\n"
     "\u2022 Concrete, end-to-end clinical use case (cataract)\n"
     "\u2022 Balanced view: candid list of challenges & enablers\n"
     "\u2022 Strong author team (NUS, SERI, Yale, Moorfields/UCL)\n"
     "\u2022 Timely topic in a high-impact Q1 venue", GREEN, body_size=16)
card(s, Inches(6.8), Inches(1.65), Inches(6.0), Inches(4.4), "Limitations",
     "\u2022 Comment article: no implementation, dataset or quantitative evaluation\n"
     "\u2022 Agent architecture, models, and communication protocol not specified\n"
     "\u2022 No latency, accuracy or cost measurements for the workflow\n"
     "\u2022 Security / privacy of inter-agent communication only briefly mentioned\n"
     "\u2022 Edge-device feasibility flagged but not analysed", ACCENT, body_size=16)
text(s, Inches(0.5), Inches(6.2), Inches(12.3), Inches(0.7),
     "Open questions: How should agents communicate securely and efficiently on resource-constrained edge "
     "devices? How do we validate and audit multi-agent decisions?", size=15, color=NAVY, italic=True,
     align=PP_ALIGN.CENTER)

# 15. Thanks -----------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, SW, SH, fill=NAVY)
box(s, 0, Inches(4.1), SW, Inches(0.07), fill=ACCENT)
text(s, 0, Inches(2.4), SW, Inches(1.2), "Thank You", size=54, color=WHITE, bold=True,
     align=PP_ALIGN.CENTER)
text(s, 0, Inches(3.35), SW, Inches(0.6), "Questions & Discussion", size=24, color=ORANGE,
     align=PP_ALIGN.CENTER)
text(s, Inches(1), Inches(4.5), SW - Inches(2), Inches(1.2),
     "Zou K, Goh JHL, Yang GD, Srinivasan S, Chen Q, Keane PA, Tham Y-C. Agentic artificial intelligence in eye care: "
     "is clinical autonomy finally within reach? Lancet Digit Health 2026; 8: 100967. doi:10.1016/j.landig.2025.100967",
     size=14, color=RGBColor(0xD6, 0xDB, 0xEA), align=PP_ALIGN.CENTER)
text(s, 0, Inches(6.2), SW, Inches(0.9), [
    ("Sheikh Sayed Bin Rahman", {"size": 18, "color": WHITE, "bold": True, "align": PP_ALIGN.CENTER, "space": 2}),
    ("PIC Lab, Kumoh National Institute of Technology", {"size": 15, "color": RGBColor(0xB0, 0xB8, 0xD0),
                                                         "align": PP_ALIGN.CENTER}),
])

NOTES = [
    # 1
    "Good afternoon everyone. My name is Sheikh Sayed Bin Rahman from the PIC Lab. Today I will present the paper "
    "\"Agentic Artificial Intelligence in Eye Care: Is Clinical Autonomy Finally Within Reach?\" by Ke Zou and "
    "colleagues, with Yih-Chung Tham as the corresponding author. It was published in The Lancet Digital Health in "
    "March 2026. The authors are from the National University of Singapore, the Singapore Eye Research Institute, "
    "Yale University, and Moorfields Eye Hospital together with University College London. The paper discusses how "
    "agentic AI, meaning multiple autonomous AI agents working together, could change clinical workflows in "
    "ophthalmology.",
    # 2
    "First, a quick word about the venue. The Lancet Digital Health is a Q1 journal across all major ranking systems. "
    "On SCImago it has an SJR of 6.838, and it is Q1 in every category it is indexed in: Health Informatics, Health "
    "Information Management, Medicine miscellaneous, and Decision Sciences miscellaneous. The most recent 2025 SJR "
    "edition also lists it as Q1. In Clarivate's Journal Citation Reports it is indexed in SCIE, with an impact factor "
    "of 24.1 for 2024 and 25.5 for 2025. That places it in about the top 2 percent of its categories, at the 98.1st "
    "percentile. Its Scopus CiteScore is 38.4, also Q1. One important note: this article is a Comment, which is a "
    "perspective paper. It proposes a conceptual framework and does not report experiments, so the methodology part "
    "of this talk is about the framework the authors propose.",
    # 3
    "Let me start with the motivation. Foundation models, meaning large language models and vision-language models, "
    "have pushed medical AI forward a lot. However, the authors point out that these models are reactive. They only "
    "produce an output when someone gives them an explicit prompt, and they are limited by task-specific training. "
    "They cannot keep track of changes in the clinical context, bring in new information on their own, adjust their "
    "reasoning, or take an action without a human. At the same time, ophthalmology faces a global shortage of "
    "specialists, while ageing populations and comorbidities make cases more complex. So the question of the paper is: "
    "can agentic AI take over routine clinical tasks, so that clinicians can focus on the final decisions?",
    # 4
    "Here are the key concepts. On the left is a plain foundation model: you give it a prompt and it gives you an "
    "output. In the middle, an AI agent is a foundation model with two additions: goal-oriented reasoning and the "
    "ability to use tools. On the right, agentic AI is a multi-agent design. Several specialised agents interact, "
    "share context, and dynamically coordinate how tasks are carried out. The authors highlight these key attributes: "
    "real-time adaptability, contextual awareness, autonomous decision making, planning and tool selection, and "
    "proactive actions, for example adjusting a treatment plan based on the patient's response. In short, agentic AI "
    "moves from passive prediction to actively supporting the clinician through the whole decision process.",
    # 5
    "This is the central figure of the paper, which I extracted directly from the PDF. It has three panels and it "
    "forms the methodology of the paper. Panel A shows today's traditional AI-assisted care model. Panel B shows the "
    "proposed care model with medical agentic AI. Panel C applies this idea to a concrete case: an eye agentic AI "
    "system for cataract disease. I will now go through each panel in detail.",
    # 6
    "Panel A is the traditional model, which is the baseline. The patient gives a history to the clinician. The "
    "clinician orders diagnostic tests. The history, examination results, and records are stored in the electronic "
    "health record. An AI model then reads this information and produces a preliminary diagnosis. Finally, the "
    "clinician makes the diagnosis and treatment plan. The key limitation is that the AI acts only when the clinician "
    "asks it to, and only for one isolated task. It has limited autonomy and needs a lot of clinician oversight, which "
    "takes up the clinician's time.",
    # 7
    "Panel B is the proposed care model with medical agentic AI. After the patient gives a history, a multi-agent "
    "system takes over most of the intermediate steps. As shown in the dashed box, the agents first devise a clinical "
    "plan, then order the tests, then produce a preliminary diagnosis, and finally generate reports. The clinician "
    "receives these outputs and focuses on the final diagnosis and management. The main idea is that routine or "
    "cognitively demanding tasks, such as history taking and summarising notes, are handled by the agents. This "
    "reduces the clinician's workload and improves efficiency, while the clinician stays in control of the final "
    "decision.",
    # 8
    "Why do the authors pick ophthalmology? There are several reasons. First, there is a global shortage of "
    "ophthalmologists while demand keeps growing. Second, ageing populations and comorbidities make care more complex "
    "and personalised. Third, eye care is highly multimodal, with fundus photos, slit-lamp images, OCT and clinical "
    "measurements, which suits a combination of vision-language models and language models. Fourth, it is "
    "protocol-driven: standard examinations like visual acuity, intraocular pressure, refraction and dilation map "
    "naturally onto agent tasks. Also, an eye-specific agent that has both ophthalmic and systemic knowledge can bridge "
    "gaps between subspecialties and use resources better.",
    # 9
    "Panel C is the cataract use case. On the left, in orange, are the patient's needs: an initial consultation, "
    "questions about surgery or medication, and chronic conditions. In the middle, in the red box, is the eye agentic "
    "AI system with four specialised agents working in sequence: the imaging agent, the diagnosis agent, the "
    "summarisation agent, and the scheduling agent. On the right, in green, is the ophthalmologist, who confirms the "
    "diagnosis, administers treatment, and performs the follow-up examination. So the agents prepare everything, and "
    "the human expert makes the final call.",
    # 10
    "Let's look at each agent. The first agent uses an LLM to collect the patient's eye complaints and visual "
    "demands, and then orders the preliminary tests by itself: visual acuity, intraocular pressure, refraction, "
    "dilation, and a fundus photo. Next, the diagnosis agent uses a vision-language model to analyse the ocular images "
    "and give a preliminary diagnosis: cataract and any other possible comorbidities, based on evidence such as lens "
    "opacity and reduced visual acuity. Then the summarisation agent, also LLM-based, compiles all diagnostic "
    "information into a report and a preliminary treatment plan, and forwards it to the clinician. Finally, if a "
    "significant cataract is detected, the scheduling agent is triggered. It books the surgery, recommends a suitable "
    "artificial lens implant, and prepares a tailored post-surgery follow-up plan. As a result, the consultation, "
    "diagnostic interpretation, and follow-up preparation are all completed before the clinician's final review.",
    # 11
    "Of course, there are serious challenges, and the authors list five. First, computational cost: multi-agent "
    "systems must store and process clinical history, images, and the dialogue between agents, which strains edge and "
    "mobile devices. Second, coordination is fragile. Without standard protocols and shared goals, agents might give "
    "inconsistent outputs or compete for resources, especially in complex workflows such as emergency triage. Third, "
    "emergent behaviour: for example, the scheduling and diagnosis agents might prioritise high-risk patients "
    "inconsistently, delaying urgent cases. Fourth, debugging and traceability: as the number of agents grows, it "
    "becomes hard to trace how a decision was made, which hurts transparency and clinician trust. Fifth, fairness, "
    "ethics, and patient safety are still largely unexplored for these systems.",
    # 12
    "To address these challenges, the authors propose four key enablers. First, human-in-the-loop designs, so that "
    "experts can guide the agents and correct errors in time. Second, standardised communication protocols and semantic "
    "planning frameworks to improve coordination and keep agent outputs aligned with clinical workflows and intent. "
    "Third, verifiable governance: monitoring the system's behaviour in real time, enforcing safety limits, and using "
    "models that predict how the system will interact with its environment. Fourth, for cross-specialty collaboration, "
    "for example between ophthalmic and systemic agents, shared causal reasoning and coherent context modelling are "
    "needed for interpretability. Underlying all of this, privacy and fairness must be built in from the start, "
    "otherwise clinicians and patients will not trust these systems.",
    # 13
    "To conclude: agentic AI moves medical AI from reactive prediction towards autonomous, interactive, and adaptive "
    "assistance. It could analyse clinical patterns, assist diagnosis, personalise treatment, and support long-term "
    "care. Ophthalmology is a strong first domain because it is multimodal and protocol-driven, and the cataract case "
    "shows how four agents could complete the consultation pipeline before the clinician reviews it. So, to answer the "
    "title question, my reading is: partially. The technology seems ready for supervised autonomy on routine tasks, but "
    "full autonomy still depends on solving safety, coordination, and trust.",
    # 14
    "Finally, here is my own critical assessment for our discussion. This slide is my view, not the authors'. On the "
    "strengths side, the paper gives a very clear conceptual comparison between traditional and agentic AI, a concrete "
    "end-to-end use case, and a balanced discussion of challenges. It also comes from a strong team and is published in "
    "a high-impact venue. On the limitations side, because it is a Comment, there is no implementation, no dataset, and "
    "no quantitative evaluation. The architecture, the specific models, and the communication protocol between agents "
    "are not specified, and there are no measurements of latency, accuracy, or cost. Security and privacy of the "
    "communication between agents, and feasibility on edge devices, are mentioned but not analysed. These open "
    "questions are relevant to our lab's work on secure, resource-constrained IoT communication.",
    # 15
    "Thank you for your attention. I'm happy to take any questions or comments, and the full citation of the paper is "
    "shown here on the slide.",
]
assert len(NOTES) == TOTAL
for slide, note in zip(prs.slides, NOTES):
    slide.notes_slide.notes_text_frame.text = note

assert len(prs.slides) == TOTAL
prs.save(OUT)
print("saved", OUT)
